"""Analyze various file types."""

import json
import os
import sys
from pathlib import Path

params = json.loads(os.environ.get("PARVA_PARAMS", "{}"))
filename = params.get("filename", "")
action = params.get("action", "extract_text")
max_chars = params.get("max_chars", 10000)
files_dir = Path(os.environ.get("PARVA_FILES_DIR", "/workspace/files"))

if not filename:
    print(json.dumps({"error": "No filename provided"}))
    sys.exit(0)

filepath = files_dir / filename
if not filepath.exists():
    # Try finding file in workspace
    workspace = Path(os.environ.get("PARVA_WORKSPACE", "/workspace"))
    filepath = workspace / filename
    if not filepath.exists():
        print(json.dumps({"error": f"File not found: {filename}"}))
        sys.exit(0)

suffix = filepath.suffix.lower()
result = {}

try:
    if suffix == ".pdf":
        from PyPDF2 import PdfReader
        reader = PdfReader(str(filepath))
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        result = {
            "type": "pdf",
            "pages": len(reader.pages),
            "text": text[:max_chars],
            "total_chars": len(text),
        }

    elif suffix in (".csv", ".tsv"):
        import csv
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f, delimiter="\t" if suffix == ".tsv" else ",")
            rows = list(reader)
        result = {
            "type": "csv",
            "rows": len(rows),
            "columns": len(rows[0]) if rows else 0,
            "headers": rows[0] if rows else [],
            "sample": rows[:10],
        }

    elif suffix in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(str(filepath), read_only=True)
        sheets = {}
        for name in wb.sheetnames:
            ws = wb[name]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 10:
                    break
                rows.append([str(c) if c is not None else "" for c in row])
            sheets[name] = {
                "rows": ws.max_row,
                "columns": ws.max_column,
                "sample": rows,
            }
        result = {"type": "excel", "sheets": sheets}

    elif suffix in (".docx",):
        from docx import Document
        doc = Document(str(filepath))
        text = "\n".join(p.text for p in doc.paragraphs)
        result = {
            "type": "docx",
            "paragraphs": len(doc.paragraphs),
            "text": text[:max_chars],
        }

    elif suffix in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"):
        from PIL import Image
        img = Image.open(filepath)
        result = {
            "type": "image",
            "format": img.format,
            "size": list(img.size),
            "mode": img.mode,
        }

    elif suffix in (".json",):
        data = json.loads(filepath.read_text())
        preview = json.dumps(data, indent=2)[:max_chars]
        result = {"type": "json", "preview": preview}

    else:
        # Try as text
        try:
            text = filepath.read_text(encoding="utf-8", errors="replace")
            result = {
                "type": "text",
                "text": text[:max_chars],
                "total_chars": len(text),
            }
        except Exception:
            result = {
                "type": "binary",
                "size": filepath.stat().st_size,
                "message": "Binary file — cannot extract text",
            }

except Exception as e:
    result = {"error": str(e)}

print(json.dumps(result, default=str))
